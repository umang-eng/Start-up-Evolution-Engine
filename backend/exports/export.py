import logging
import os
from typing import Any
import uuid

try:
    import pptx
    from pptx import Presentation
except ImportError:
    pptx = None

try:
    import weasyprint
except ImportError:
    weasyprint = None

logger = logging.getLogger("app.exports")

class ExportCompiler:
    def compile_blueprint_to_html(self, blueprint_data: dict[str, Any]) -> str:
        title = blueprint_data.get("executive_summary", {}).get("startup_name") or "Startup Blueprint"
        industry = blueprint_data.get("dna", {}).get("category") or "Technology"
        
        html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{title} - Startup Blueprint</title>
    <style>
        body {{ font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif; color: #1a202c; line-height: 1.6; margin: 0; padding: 0; background: #fff; }}
        @page {{ size: A4; margin: 20mm; @bottom-right {{ content: counter(page); font-size: 9pt; color: #8892b0; }} }}
        .cover-page {{ height: 100vh; display: flex; flex-direction: column; justify-content: center; align-items: center; text-align: center; padding: 40px; page-break-after: always; }}
        .cover-title {{ font-size: 42pt; font-weight: 800; color: #0f172a; margin-bottom: 10px; }}
        .cover-subtitle {{ font-size: 18pt; color: #475569; margin-bottom: 40px; }}
        .section {{ margin-bottom: 40px; page-break-inside: avoid; }}
        h1 {{ font-size: 24pt; color: #1e3a8a; border-bottom: 2px solid #e2e8f0; padding-bottom: 8px; margin-top: 40px; }}
        h2 {{ font-size: 16pt; color: #2563eb; margin-top: 25px; }}
        .card {{ background: #ffffff; border: 1px solid #e2e8f0; border-radius: 12px; padding: 20px; margin-bottom: 20px; }}
        table {{ width: 100%; border-collapse: collapse; margin: 20px 0; }}
        th, td {{ padding: 12px; border-bottom: 1px solid #e2e8f0; text-align: left; }}
        th {{ background-color: #f8fafc; font-weight: 600; }}
        .metric-grid {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 20px; margin: 20px 0; }}
        .metric-card {{ background: #eff6ff; border: 1px solid #bfdbfe; border-radius: 8px; padding: 15px; text-align: center; }}
        .metric-value {{ font-size: 20pt; font-weight: 700; color: #1e3a8a; }}
    </style>
</head>
<body>
    <div class="cover-page">
        <div class="cover-title">{title}</div>
        <div class="cover-subtitle">Strategic & Operational Blueprint</div>
        <div style="margin-top: 50px; color: #64748b;">Industry: {industry}</div>
    </div>

    <div class="section">
        <h1>1. Executive Summary</h1>
        <div class="card">
            <h2>Vision</h2>
            <p>{blueprint_data.get("executive_summary", {}).get("vision", "Not available")}</p>
            <h2>Summary</h2>
            <p>{blueprint_data.get("executive_summary", {}).get("summary", "Not available")}</p>
        </div>
    </div>
"""
        if "dna" in blueprint_data:
            dna = blueprint_data["dna"]
            html += f"""
    <div class="section page-break">
        <h1>2. Startup DNA</h1>
        <div class="card">
            <h2>Business Model</h2>
            <p>{dna.get("business_model", "Not defined")}</p>
            <h2>Target Market</h2>
            <p>{dna.get("target_market", "Not defined")}</p>
        </div>
    </div>
"""
        if "features" in blueprint_data:
            feats = blueprint_data["features"].get("features", [])
            html += """
    <div class="section page-break">
        <h1>3. Product Features</h1>
        <table>
            <thead><tr><th>Feature</th><th>Priority</th></tr></thead>
            <tbody>
            """
            for feat in feats:
                html += f"""<tr><td><strong>{feat.get("name")}</strong><br><small>{feat.get("description")}</small></td><td>{feat.get("priority")}</td></tr>"""
            html += "</tbody></table></div>"

        if "team" in blueprint_data:
            roles = blueprint_data["team"].get("org_chart", [])
            html += """
    <div class="section page-break">
        <h1>4. Organizational Structure</h1>
        <table>
            <thead><tr><th>Role</th><th>Department</th><th>Salary (USD)</th></tr></thead>
            <tbody>
            """
            for role in roles:
                html += f"""<tr><td><strong>{role.get("title")}</strong></td><td>{role.get("department")}</td><td>${role.get("estimated_salary_usd", 0):,}</td></tr>"""
            html += "</tbody></table></div>"

        if "swot" in blueprint_data:
            swot = blueprint_data["swot"]
            html += """
    <div class="section page-break">
        <h1>5. SWOT Analysis</h1>
        <div style="display: flex; gap: 20px;">
            <div class="card" style="flex: 1;"><h3>Strengths</h3><ul>
            """
            for s in swot.get("strengths", []): html += f"<li>{s}</li>"
            html += """</ul></div><div class="card" style="flex: 1;"><h3>Weaknesses</h3><ul>"""
            for w in swot.get("weaknesses", []): html += f"<li>{w}</li>"
            html += """</ul></div></div><div style="display: flex; gap: 20px;"><div class="card" style="flex: 1;"><h3>Opportunities</h3><ul>"""
            for o in swot.get("opportunities", []): html += f"<li>{o}</li>"
            html += """</ul></div><div class="card" style="flex: 1;"><h3>Threats</h3><ul>"""
            for t in swot.get("threats", []): html += f"<li>{t}</li>"
            html += "</ul></div></div></div>"

        if "cost" in blueprint_data:
            cost = blueprint_data["cost"]
            funding_req = cost.get("funding_requirements", {})
            opt = funding_req.get("optimal_target_usd", 0)
            mvp = cost.get("mvp_cost_estimate", 0)
            yr1 = cost.get("year_1_cost_estimate", 0)
            html += f"""
    <div class="section page-break">
        <h1>6. Financial Projections</h1>
        <div class="metric-grid">
            <div class="metric-card"><div class="metric-value">${opt:,}</div><div>Optimal Target</div></div>
            <div class="metric-card"><div class="metric-value">${mvp:,}</div><div>MVP Cost</div></div>
            <div class="metric-card"><div class="metric-value">${yr1:,}</div><div>Year 1 Cost</div></div>
        </div>
    </div>
"""
        html += "</body></html>"
        return html

    def compile_blueprint_to_pdf(self, blueprint_data: dict[str, Any]) -> bytes | None:
        if not weasyprint:
            return None
        try:
            html_content = self.compile_blueprint_to_html(blueprint_data)
            return weasyprint.HTML(string=html_content).write_pdf()
        except Exception as e:
            logger.error("Failed to compile blueprint PDF", exc_info=e)
            return None

    def compile_blueprint_to_deck(self, blueprint_data: dict[str, Any]) -> bytes | None:
        if not pptx:
            return None
        try:
            prs = Presentation()
            title_layout = prs.slide_layouts[0]
            bullet_layout = prs.slide_layouts[1]

            slide1 = prs.slides.add_slide(title_layout)
            slide1.shapes.title.text = blueprint_data.get("executive_summary", {}).get("startup_name") or "Startup Pitch Deck"
            slide1.placeholders[1].text = "Strategic AI Platform Blueprint"

            if "dna" in blueprint_data:
                slide2 = prs.slides.add_slide(bullet_layout)
                slide2.shapes.title.text = "Value Proposition"
                tf = slide2.shapes.placeholders[1].text_frame
                tf.text = blueprint_data["dna"].get("business_model", "Not Defined")[:200]

            temp_path = f"/tmp/deck-{uuid.uuid4()}.pptx"
            prs.save(temp_path)
            with open(temp_path, "rb") as f:
                pptx_bytes = f.read()
            os.remove(temp_path)
            return pptx_bytes
        except Exception as e:
            logger.error("Failed to compile PowerPoint deck", exc_info=e)
            return None

export_compiler = ExportCompiler()
